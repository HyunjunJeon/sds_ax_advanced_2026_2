"""새 번호 실습의 결과물 연결과 중요한 실패 경계를 검사한다. 모든 검수·모델은 테스트용 가상 자료다."""

from copy import deepcopy
import csv
import importlib.util
import json
from pathlib import Path

from openpyxl import load_workbook
import pytest

from business_lab.agents import COMMON_PROMPT
from business_lab.authoring import draft_cases, freeze, load_frozen, read_workbook, write_workbook
from business_lab.dataset import load_fixtures
from business_lab.evidence import calibrate_partitions, validate_plan, write_final_report
from business_lab.lesson_runs import compare_scored, grade_recording, read_recording, record_release
import learner_evaluator

ROOT = Path(__file__).resolve().parents[1]


def lesson(filename, **settings):
    spec = importlib.util.spec_from_file_location("new_lesson_" + filename.replace(".", "_"), ROOT / filename)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    for key, value in settings.items():
        if hasattr(module, key):
            setattr(module, key, value)
    return module


def approved_cards():
    cards = draft_cases()
    for card in cards:
        card["metadata"].update(review_status="approved", reviewer="test-reviewer", review_reason="자동 테스트의 가상 검수 자료")
    return cards


@pytest.fixture
def frozen(tmp_path):
    return freeze(approved_cards(), load_fixtures(), tmp_path / "dataset.json",
                  scenario_ids=["normal", "timeout-after", "clarification-episode"])


def record(frozen, tmp_path, release="baseline"):
    return record_release(frozen, release=release, prompt=COMMON_PROMPT + ("\n테스트용 조건" if release == "candidate" else ""),
                          repeats=2, mode="scripted", use_langfuse=False, out=tmp_path / f"{release}.json")


def test_reviewed_workbook_is_the_actual_execution_input(tmp_path):
    path = tmp_path / "draft.xlsx"
    write_workbook(draft_cases(), path)
    with pytest.raises(ValueError, match="검수"):
        freeze(read_workbook(path), load_fixtures(), tmp_path / "blocked.json", scenario_ids=["normal"])
    assert not (tmp_path / "blocked.json").exists()
    book = load_workbook(path)
    sheet = book["시나리오"]
    sheet.cell(2, 2).value = json.dumps({"message": "테스트에서 검수한 새 입력", "order_id": "order-a", "request_id": "test-review-id"})
    for column, value in ((15, "approved"), (16, "test-reviewer"), (17, "정책과 초기 상태를 대조한 테스트 자료")):
        sheet.cell(2, column).value = value
    book.save(path)
    data = freeze(read_workbook(path), load_fixtures(), tmp_path / "dataset.json", scenario_ids=["normal"])
    raw = record(data, tmp_path)
    assert raw["rows"][0]["artifact"]["initial_request"]["message"] == "테스트에서 검수한 새 입력"
    assert len(raw["rows"]) == 2
    assert all(row["verdicts"] == [] for row in raw["rows"])


@pytest.mark.parametrize("mutation,match", [("empty_gold", "validation"), ("no_reviewer", "검수"),
                                           ("duplicate", "중복"), ("family_leak", "family_id")])
def test_invalid_review_cannot_become_a_golden(tmp_path, mutation, match):
    cards = approved_cards()
    if mutation == "empty_gold":
        cards[0]["expected_output"]["final_state"] = {}
    elif mutation == "no_reviewer":
        cards[0]["metadata"]["reviewer"] = ""
    elif mutation == "duplicate":
        extra = deepcopy(cards[0])
        extra["metadata"]["scenario_id"] = "different-id-same-case"
        extra["input"]["request_id"] = "another-request"
        cards.append(extra)
    else:
        cards[0]["metadata"]["family_id"] = "incident-resume-pending"
    with pytest.raises(ValueError, match=match):
        freeze(cards, load_fixtures(), tmp_path / "invalid.json")


def test_dataset_and_raw_recording_edits_are_detected(frozen, tmp_path):
    path = tmp_path / "dataset.json"
    data = json.loads(path.read_text())
    data["snapshot"]["cards"][0]["input"]["message"] = "changed after freeze"
    path.write_text(json.dumps(data))
    with pytest.raises(ValueError, match="변경"):
        load_frozen(path)
    record(frozen, tmp_path)
    raw = json.loads((tmp_path / "baseline.json").read_text())
    raw["rows"].pop()
    (tmp_path / "baseline.json").write_text(json.dumps(raw))
    with pytest.raises(ValueError, match="변경"):
        read_recording(tmp_path / "baseline.json")


def test_episode_question_completion_and_fresh_trials_are_distinct(frozen, tmp_path):
    raw = record(frozen, tmp_path)
    episodes = [r["artifact"] for r in raw["rows"] if r["artifact"]["scenario_id"] == "clarification-episode"]
    for a in episodes:
        assert len(a["turns"]) == 2
        first, last = a["turns"]
        assert first["response"]["kind"] == "clarification" and first["event_ids"] == []
        assert first["final_state"]["orders"]["order-a"]["status"] == "paid"
        assert last["response"]["kind"] == "completed"
        assert last["initial_state"] == first["final_state"]
        assert a["final_state"]["orders"]["order-a"]["status"] == "cancelled"
    assert episodes[0]["initial_state"] == episodes[1]["initial_state"]
    assert episodes[0]["episode_id"] != episodes[1]["episode_id"]
    graded = grade_recording(raw)
    assert all(next(v for v in r["verdicts"] if v["name"] == "episode")["status"] == "PASS" for r in graded["rows"])


def test_early_stop_is_not_a_successful_episode(frozen, tmp_path):
    raw = record(frozen, tmp_path)
    row = next(r for r in raw["rows"] if r["artifact"]["scenario_id"] == "clarification-episode")
    row["artifact"]["turns"] = row["artifact"]["turns"][:1]
    report = grade_recording(raw)
    row = next(r for r in report["rows"] if r["artifact"]["scenario_id"] == "clarification-episode")
    assert next(v for v in row["verdicts"] if v["name"] == "episode")["status"] == "FAIL"


def test_missing_trial_stays_in_denominator_and_not_called_regression(frozen, tmp_path):
    b, c = (grade_recording(record(frozen, tmp_path, release)) for release in ("baseline", "candidate"))
    c["rows"].pop()
    report = compare_scored(b, c)
    summary = report["comparison"]["summaries"]["candidate"]
    assert summary["planned"] == 6 and summary["found"] == 5 and summary["statuses"]["MISSING"] == 1
    assert report["comparison"]["all_repeat_transitions"]["uncomparable"] == 1
    assert report["comparison"]["decision"] == "HOLD"


def test_changed_protocol_or_judge_cannot_be_compared(frozen, tmp_path):
    b, c = (grade_recording(record(frozen, tmp_path, release)) for release in ("baseline", "candidate"))
    changed = deepcopy(c)
    changed["config"]["max_model_calls"] += 1
    with pytest.raises(ValueError, match="조건"):
        compare_scored(b, changed)
    changed = deepcopy(c)
    changed["config"]["judge_model"] = "different-judge"
    with pytest.raises(ValueError, match="재채점"):
        compare_scored(b, changed)


def test_human_review_needs_both_classes_in_separate_partitions(frozen, tmp_path):
    scored = grade_recording(record(frozen, tmp_path))
    labels = []
    for row in scored["rows"]:
        row["verdicts"].append({"name": "response_semantics", "status": "PASS"})
        labels.append({"artifact_id": row["artifact"]["artifact_id"],
                       "rubric_version": scored["config"]["response_rubric_version"],
                       "human_verdict": "PASS", "reviewer": "test-human", "reason": "가상 검수"})
    result = calibrate_partitions(scored, labels)
    assert result["status"] == "insufficient_coverage"
    assert len(result["coverage_gaps"]) == 2
    for row, label in zip(scored["rows"], labels):
        if row["artifact"]["trial"] == 2:
            label["human_verdict"] = "FAIL"
    result = calibrate_partitions(scored, labels)
    assert result["status"] == "reviewed" and result["coverage_gaps"] == []
    assert result["total"]["false_pass"] == 3
    comparison = {"evaluation_run_ids": [scored["run_id"]], "snapshot_hash": "test", "source_runs": [],
                  "comparison": {"decision_reasons": [], "summaries": {}, "pairs": []}}
    plan = {field: "test" for field in ("fact", "hypothesis", "alternative", "change", "success_criteria", "owner")}
    path = tmp_path / "report.md"
    write_final_report(comparison, plan, path, calibration=result,
                       holdout={"comparison": {"decision": "READY_FOR_HUMAN_REVIEW"}})
    assert "HOLD" in path.read_text() and "FAIL을 PASS" in path.read_text()


def test_all_numbered_steps_connect_without_fabricating_real_reviews(tmp_path, monkeypatch):
    settings = {"WORK": tmp_path, "MODE": "scripted", "USE_LANGFUSE": False, "USE_JUDGE": False,
        "FETCH_REMOTE_OBSERVATIONS": False, "REPEATS": 2, "SCENARIO_IDS": ["normal", "timeout-after", "clarification-episode"]}
    filenames = {"DRAFT_XLSX": "01.xlsx", "DATASET": "02.json", "BASELINE": "03.json", "CHANGE_PLAN": "04.json",
                 "SCORED_BASELINE": "05.json", "HUMAN_LABELS": "06.csv", "CALIBRATION": "06.json", "CANDIDATE": "candidate.json",
                 "COMPARISON": "07.json", "REGRESSION_DRAFT": "08.xlsx", "HOLDOUT": "holdout.json", "FINAL_REPORT": "09.md"}
    settings.update({key: tmp_path / name for key, name in filenames.items()})
    lesson("01_draft_scenarios.py", **settings).main()
    second = lesson("02_freeze_dataset.py", **settings)
    with pytest.raises(ValueError, match="검수"):
        second.main()
    book = load_workbook(settings["DRAFT_XLSX"])
    for row in book["시나리오"].iter_rows(min_row=2):
        row[14].value, row[15].value, row[16].value = "approved", "test-reviewer", "자동 테스트 전용 자료"
    book.save(settings["DRAFT_XLSX"])
    second.main()
    lesson("03_record_baseline.py", **settings).main()
    lesson("04_inspect_traces.py", **settings).main()
    raw = read_recording(settings["BASELINE"])
    with pytest.raises(ValueError, match="직접"):
        validate_plan(settings["CHANGE_PLAN"], raw)
    plan = json.loads(settings["CHANGE_PLAN"].read_text())
    plan.update(artifact_id=raw["rows"][0]["artifact"]["artifact_id"], evidence_ids=["final_state"],
                fact="테스트 관찰", hypothesis="테스트 가설", alternative="테스트 대안", change="테스트 한 조건",
                success_criteria="테스트 판정 기준", owner="test-author")
    settings["CHANGE_PLAN"].write_text(json.dumps(plan))
    lesson("05_evaluate_runs.py", **settings).main()
    review = lesson("06_review_judge.py", **settings)
    review.main()
    assert not settings["CALIBRATION"].exists()  # 검수자를 생성해 채우지 않는다.
    with settings["HUMAN_LABELS"].open(encoding="utf-8-sig", newline="") as f:
        labels = list(csv.DictReader(f))
    assert all(not row["human_verdict"] and not row["reviewer"] for row in labels)
    for row in labels:
        row.update(human_verdict="PASS", reviewer="test-human", reason="테스트 판정")
    with settings["HUMAN_LABELS"].open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(labels[0]))
        writer.writeheader(); writer.writerows(labels)
    review.main()
    seventh = lesson("07_compare_candidate.py", **settings)
    with pytest.raises(ValueError, match="baseline과 같"):
        seventh.main()
    monkeypatch.setattr(seventh.candidate_agent, "PROMPT", COMMON_PROMPT + "\n테스트용 조건")
    seventh.main()
    eighth = lesson("08_promote_regression.py", **settings)
    with pytest.raises(ValueError, match="비식별화"):
        eighth.main()
    eighth.SANITIZATION_RECORD = "테스트용 가상 데이터만 사용"
    eighth.main()
    draft = read_workbook(settings["REGRESSION_DRAFT"])[0]
    assert draft["metadata"]["review_status"] == "pending" and draft["expected_output"]["final_state"] == {}
    assert draft["metadata"]["source_artifact_id"] == plan["artifact_id"]
    eighth.ACTION = "holdout"
    eighth.main()
    lesson("09_report_decision.py", **settings).main()
    assert "HOLD" in settings["FINAL_REPORT"].read_text()
    assert "scripted" in settings["FINAL_REPORT"].read_text()
